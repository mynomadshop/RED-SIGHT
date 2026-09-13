"""Bounded document and tabular operations for the bundled procedural skills."""

from __future__ import annotations

import csv
import hashlib
import itertools
import json
import os
import re
import tempfile
from collections import Counter
from contextlib import contextmanager
from pathlib import Path

TOOL_SPECS = {
    "documents.extract": {
        "description": "Extract text from PDF, DOCX or UTF-8 text with explicit truncation; never run document macros.",
        "risk": "read", "approval": False, "agent": True,
        "params": "path:str, max_chars?:int<=100000, max_pages?:int<=100",
    },
    "data.profile": {
        "description": "Inspect CSV, TSV or XLSX headers, sample rows, blanks and exact duplicate rows with bounded memory.",
        "risk": "read", "approval": False, "agent": True,
        "params": "path:str, sheet?:str, max_rows?:int<=100000",
    },
    "data.merge": {
        "description": "Merge CSV/TSV files with matching headers to a NEW CSV. Preserve text IDs; optionally remove exact duplicates. Refuse existing outputs and partial results.",
        "risk": "write", "approval": True, "agent": True,
        "params": "paths:list[str], output_path:str, deduplicate?:bool, max_rows?:int<=1000000",
    },
}


def _limit(params: dict, key: str, default: int, maximum: int) -> int:
    return max(1, min(maximum, int(params.get(key, default))))


def _source(raw: str, validate) -> Path:
    path = validate(raw)
    if not path.is_file():
        raise FileNotFoundError("Input file does not exist")
    if path.stat().st_size > 250 * 1024 * 1024:
        raise ValueError("Input exceeds the 250 MB per-file limit; split the data first")
    return path


def _headers(values) -> list[str]:
    headers = [str("" if value is None else value).strip() for value in values]
    if not headers or any(not value for value in headers) or len(headers) != len(set(headers)):
        raise ValueError("Headers must be nonempty and unique; inspect and repair the source first")
    if len(headers) > 500:
        raise ValueError("Input exceeds 500 columns")
    return headers


def _spreadsheet_cell(value: str) -> tuple[str, bool]:
    # Keep signed numeric values (including scientific notation) unchanged.
    # Formula-like text, including column titles, must stay literal in Excel.
    signed_number = re.fullmatch(r"[+-](?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?", value.strip())
    if value.lstrip().startswith(("=", "+", "-", "@")) and not signed_number:
        return "'" + value, True
    return value, False


@contextmanager
def _table(path: Path, sheet: str = ""):
    if path.suffix.lower() in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
            yield _headers(next(reader, [])), reader
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
            rows = worksheet.iter_rows(values_only=True)
            yield _headers(next(rows, [])), rows
        finally:
            workbook.close()
    else:
        raise ValueError("Use a UTF-8 CSV, TSV or XLSX file")


def profile(params: dict, validate) -> dict:
    path = _source(str(params.get("path", "")), validate)
    limit = _limit(params, "max_rows", 10000, 100000)
    samples = []
    blanks: Counter = Counter()
    seen: set[bytes] = set()
    count = duplicates = malformed = 0
    truncated = False
    with _table(path, str(params.get("sheet", ""))) as (headers, rows):
        for index, raw in enumerate(rows):
            if index >= limit:
                truncated = True
                break
            count += 1
            values = ["" if item is None else str(item) for item in raw]
            if len(values) != len(headers):
                malformed += 1
            record = dict(zip(headers, values, strict=False))
            for column in headers:
                if not record.get(column, "").strip():
                    blanks[column] += 1
            digest = hashlib.sha256(json.dumps(values, ensure_ascii=False).encode()).digest()
            duplicates += digest in seen
            seen.add(digest)
            if len(samples) < 5:
                samples.append({key: value[:500] for key, value in record.items()})
    return {"ok": True, "path": str(path), "columns": headers, "rows_scanned": count,
            "truncated": truncated, "duplicate_rows": duplicates, "malformed_rows": malformed,
            "blank_cells": {column: blanks[column] for column in headers}, "sample_rows": samples,
            "note": "Counts apply to scanned rows. XLSX formulas are shown as formulas, not recalculated."}


def merge(params: dict, validate) -> dict:
    raw_paths = params.get("paths")
    if not isinstance(raw_paths, list) or not 1 <= len(raw_paths) <= 100:
        raise ValueError("Provide between 1 and 100 CSV/TSV paths")
    paths = [_source(str(raw), validate) for raw in raw_paths]
    if any(path.suffix.lower() not in {".csv", ".tsv"} for path in paths):
        raise ValueError("Merge accepts CSV/TSV inputs; profile XLSX sheets first and export the intended sheet")
    output = validate(str(params.get("output_path", "")), write=True)
    if output.suffix.lower() != ".csv":
        raise ValueError("The output path must end in .csv")
    if output.exists() or output.resolve() in {path.resolve() for path in paths}:
        raise FileExistsError("Choose a new output path; existing files are never overwritten")
    if not output.parent.is_dir():
        raise FileNotFoundError("The output directory must already exist")
    deduplicate = params.get("deduplicate", False)
    if not isinstance(deduplicate, bool):
        raise ValueError("deduplicate must be true or false")
    limit = _limit(params, "max_rows", 250000, 1000000)
    total = written = removed = escaped = 0
    seen: set[bytes] = set()
    headers = None
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8-sig", newline="",
                                         dir=output.parent, delete=False) as handle:
            temporary = Path(handle.name)
            writer = csv.writer(handle)
            for path in paths:
                with _table(path) as (current_headers, rows):
                    if headers is None:
                        headers = current_headers
                        safe_headers = [_spreadsheet_cell(value) for value in headers]
                        writer.writerow([value for value, _ in safe_headers])
                        escaped += sum(changed for _, changed in safe_headers)
                    if current_headers != headers:
                        raise ValueError("Input headers/order differ; no output was published")
                    for row in rows:
                        total += 1
                        if total > limit:
                            raise ValueError("Row limit exceeded; no partial output was published")
                        if len(row) != len(headers):
                            raise ValueError("Malformed row; no output was published")
                        digest = hashlib.sha256(json.dumps(row, ensure_ascii=False).encode()).digest()
                        if deduplicate and digest in seen:
                            removed += 1
                            continue
                        if deduplicate:
                            seen.add(digest)
                        safe = []
                        for value in row:
                            value, changed = _spreadsheet_cell(value)
                            escaped += changed
                            safe.append(value)
                        writer.writerow(safe)
                        written += 1
            handle.flush()
            os.fsync(handle.fileno())
        # Hard-link publication is atomic and refuses a path created since the
        # preflight check. Supported by NTFS and normal Linux filesystems.
        os.link(temporary, output)
        return {"ok": True, "path": str(output), "input_rows": total, "output_rows": written,
                "duplicates_removed": removed, "formula_like_cells_escaped": escaped,
                "source_files": [str(path) for path in paths], "columns": headers}
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def extract(params: dict, validate) -> dict:
    path = _source(str(params.get("path", "")), validate)
    limit = _limit(params, "max_chars", 30000, 100000)
    page_limit = _limit(params, "max_pages", 25, 100)
    parts = []
    size = 0
    truncated = False
    pages = None

    def append(text: str) -> bool:
        nonlocal size, truncated
        available = limit - size
        parts.append(text[:available])
        size += min(len(text), available)
        if len(text) > available:
            truncated = True
        return size < limit

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise ValueError("The PDF is password protected")
            pages = len(document)
            for index in range(min(pages, page_limit)):
                if not append(f"\n[Page {index + 1}]\n" + document[index].get_text()):
                    truncated = True
                    break
            truncated = truncated or pages > page_limit
    elif suffix == ".docx":
        from docx import Document

        document = Document(path)
        paragraphs = (paragraph.text + "\n" for paragraph in document.paragraphs)
        tables = ("\t".join(cell.text for cell in row.cells) + "\n"
                  for table in document.tables for row in table.rows)
        for text in itertools.chain(paragraphs, tables):
            if not append(text):
                truncated = True
                break
    elif suffix in {".txt", ".md", ".json", ".log", ".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", errors="replace") as handle:
            append(handle.read(limit + 1))
    else:
        raise ValueError("Supported formats: PDF, DOCX, TXT, Markdown, JSON, logs, CSV and TSV")
    text = "".join(parts)
    return {"ok": True, "path": str(path), "text": text, "truncated": truncated, "pages": pages,
            "note": "Text extraction only; scanned images need OCR. Tables may require layout review."}


def execute(tool: str, params: dict, validate) -> dict:
    return {"data.profile": profile, "data.merge": merge, "documents.extract": extract}[tool](params, validate)
